"""
Totumseg - Jet x Mapfre xls
---------------------------
Le as planilhas (.xlsm/.xlsx) da pasta base, escolhe o modelo .xls da Mapfre
pelo nome do arquivo (tipo do imovel + valor), cola os dados no modelo e salva
na pasta destino com o mesmo nome do arquivo original (extensao .xls).

Dependencias (requirements.txt):
    customtkinter
    pywin32
    openpyxl
"""
import datetime
import glob
import json
import os
import queue
import threading
import time
from tkinter import filedialog, messagebox

import customtkinter as ctk

try:
    import pythoncom
    import win32com.client as win32
    TEM_EXCEL_COM = True
except ImportError:  # fora do Windows / sem pywin32
    TEM_EXCEL_COM = False

NOME_APP = "Totumseg - Jet x Mapfre xls"

# --------------------------------------------------------------------------
# Modelos aceitos (os arquivos precisam ter exatamente estes nomes + .xls)
# --------------------------------------------------------------------------
MODELOS_POR_RAMO = {
    "APTO":         ["200", "300", "400"],
    "CASA_MADEIRA": ["200", "300", "400"],
    "CASA":         ["200", "300", "400"],
    "COMER":        ["200", "400", "600"],
    "ESCRIT":       ["200", "300", "600"],
}
MODELOS_VALIDOS = [f"{r}_{v}" for r, vs in MODELOS_POR_RAMO.items() for v in vs]

# --------------------------------------------------------------------------
# Paleta: predominantemente branco, com azul de destaque
# --------------------------------------------------------------------------
BRANCO = "#FFFFFF"
FUNDO_SUAVE = "#F8FAFC"
AZUL = "#2563EB"
AZUL_ESCURO = "#1D4ED8"
AZUL_CLARO = "#EFF6FF"
BORDA = "#E2E8F0"
TEXTO = "#0F172A"
TEXTO_SUAVE = "#64748B"
VERDE = "#15803D"
VERMELHO = "#B91C1C"
LARANJA = "#B45309"

ARQ_CONFIG = os.path.join(os.path.expanduser("~"), ".totumseg_jet_mapfre.json")

ctk.set_appearance_mode("light")


# --------------------------------------------------------------------------
# Regras de negocio (mesma logica do programa original)
# --------------------------------------------------------------------------
def identificar_modelo(nome_arquivo):
    """Devolve 'RAMO_VALOR' (ex.: 'APTO_300') ou None se nao identificar."""
    nome = nome_arquivo.upper()
    ramo = ""
    valor = ""

    if "MADEIRA" in nome:
        ramo = "CASA_MADEIRA"
    elif "CASA" in nome:
        ramo = "CASA"
    elif "APART" in nome or "APTO" in nome:
        ramo = "APTO"
    elif "COMER" in nome:
        ramo = "COMER"
    elif "ESCRIT" in nome:
        ramo = "ESCRIT"

    for v in ("200", "300", "400", "600"):
        if f"{v}000" in nome:
            valor = v
            break

    if ramo and valor:
        return f"{ramo}_{valor}"
    return None


def sanitizar(valor):
    """Ajustes que a Mapfre exige (pais e mascara de CPF/CNPJ)."""
    if isinstance(valor, str):
        texto = valor.strip()
        if texto.upper() == "BRASIL":
            return "1-BRASIL"
        if len(texto) == 11 and texto.isdigit():
            return f"{texto[:3]}.{texto[3:6]}.{texto[6:9]}-{texto[9:]}"
        if len(texto) == 14 and texto.isdigit():
            return f"{texto[:2]}.{texto[2:5]}.{texto[5:8]}/{texto[8:12]}-{texto[12:]}"
    return valor


def encurtar(texto, limite=58):
    return texto if len(texto) <= limite else "…" + texto[-(limite - 1):]


# --------------------------------------------------------------------------
# Motor de processamento
#
# Tres ganhos de velocidade em relacao a versao anterior:
#  1. A planilha de origem e' lida direto do arquivo (sem abrir no Excel).
#     Se algo nela for incomum (formulas, erros, horas), cai no Excel sozinho.
#  2. O modelo .xls (3 MB) fica aberto e e' reaproveitado entre arquivos do
#     mesmo modelo: grava -> salva copia -> restaura -> confere. Se a conferencia
#     falhar, o modelo e' descartado e reaberto limpo.
#  3. Varios Excel trabalham em paralelo.
# --------------------------------------------------------------------------
COL_INI, COL_FIM, LINHA_INI = 2, 40, 6   # B..AN, a partir da linha 6
LARGURA = COL_FIM - COL_INI + 1
TAM_BLOCO = 500                          # linhas lidas por chamada (leitura via Excel)
TAM_SNAP = 40                            # linhas do modelo guardadas para restaurar
MAX_EXCEL = 6                            # maximo de Excel abertos em paralelo


def col_letra(n):
    texto = ""
    while n:
        n, resto = divmod(n - 1, 26)
        texto = chr(65 + resto) + texto
    return texto


L_INI, L_FIM = col_letra(COL_INI), col_letra(COL_FIM)


def _bool(x):
    return None if x is None else bool(x)


def _resumir_erro(e):
    texto = str(e).strip().splitlines()[0] if str(e).strip() else type(e).__name__
    return texto if len(texto) <= 110 else texto[:107] + "…"


# ---------------- leitura da origem ----------------
class _PrecisaDoExcel(Exception):
    pass


_ERROS_EXCEL = {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def ler_origem_rapido(caminho):
    """Le direto do arquivo (primeira aba, a partir da linha 6) sem abrir o Excel."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise _PrecisaDoExcel()

    wb = load_workbook(caminho, read_only=True, data_only=False, keep_links=False)
    try:
        ws = wb.worksheets[0]
        linhas = []
        for row in ws.iter_rows(min_row=LINHA_INI, min_col=COL_INI, max_col=COL_FIM, values_only=True):
            row = tuple(row) + (None,) * (LARGURA - len(row))
            v = row[0]
            if v is None or str(v).strip() == "":
                break
            for x in row:
                if isinstance(x, str):
                    if x.startswith("=") or x in _ERROS_EXCEL:
                        raise _PrecisaDoExcel()
                elif x is not None and not isinstance(x, (int, float, bool, datetime.datetime)):
                    raise _PrecisaDoExcel()
            linhas.append(row)
        return linhas
    finally:
        wb.close()


def ler_origem_excel(ws):
    """Le em blocos pelo Excel; para na primeira linha com B vazio."""
    linhas = []
    r = LINHA_INI
    while r < 1_048_576:
        bloco = ws.Range(f"{L_INI}{r}:{L_FIM}{r + TAM_BLOCO - 1}").Value
        for lin in bloco:
            v = lin[0]
            if v is None or str(v).strip() == "":
                return linhas
            linhas.append(lin)
        r += TAM_BLOCO
    return linhas


def obter_linhas(excel, caminho):
    try:
        return ler_origem_rapido(caminho)
    except Exception:
        pass  # qualquer duvida: o Excel le (e' o caminho original)
    wb = excel.Workbooks.Open(os.path.abspath(caminho), ReadOnly=True, UpdateLinks=False)
    try:
        return ler_origem_excel(wb.Sheets(1))
    finally:
        wb.Close(SaveChanges=False)


# ---------------- gravacao no modelo ----------------
def agrupar_trechos(indices):
    """[0,1,2,5,6] -> [(0,2),(5,6)]"""
    trechos = []
    for i in indices:
        if trechos and i == trechos[-1][1] + 1:
            trechos[-1][1] = i
        else:
            trechos.append([i, i])
    return [(a, b) for a, b in trechos]


def gravar_modelo(ws, linhas):
    """
    Grava os valores da origem no modelo, so em celulas destravadas e so onde a
    origem tem valor (igual ao programa original), mas em blocos por coluna.
    Devolve (trechos_gravados, bloco_inteiro_destravado).
    """
    ultima = LINHA_INI + len(linhas) - 1
    bloco = ws.Range(f"{L_INI}{LINHA_INI}:{L_FIM}{ultima}")
    bloco_destravado = _bool(bloco.Locked) is False
    escritos = []

    for j in range(LARGURA):
        letra = col_letra(COL_INI + j)
        novos = [sanitizar(lin[j]) for lin in linhas]
        indices = [i for i, v in enumerate(novos) if v is not None]
        if not indices:
            continue

        travada = False if bloco_destravado else _bool(
            ws.Range(f"{letra}{LINHA_INI}:{letra}{ultima}").Locked)

        if travada is True:
            continue
        if travada is False:
            for ini, fim in agrupar_trechos(indices):
                r1, r2 = LINHA_INI + ini, LINHA_INI + fim
                if ini == fim:
                    ws.Range(f"{letra}{r1}").Value = novos[ini]
                else:
                    ws.Range(f"{letra}{r1}:{letra}{r2}").Value = tuple((novos[i],) for i in range(ini, fim + 1))
                escritos.append((j, r1, r2))
        else:  # coluna com celulas travadas e destravadas misturadas: confere uma a uma
            for i in indices:
                cel = ws.Cells(LINHA_INI + i, COL_INI + j)
                if not cel.Locked:
                    cel.Value = novos[i]
                    escritos.append((j, LINHA_INI + i, LINHA_INI + i))
    return escritos, bloco_destravado


class ModeloAberto:
    """Modelo .xls aberto (somente leitura) e reaproveitavel entre arquivos."""

    def __init__(self, excel, caminho, chave):
        self.chave = chave
        self.wb = excel.Workbooks.Open(os.path.abspath(caminho), ReadOnly=True, UpdateLinks=False,
                                       IgnoreReadOnlyRecommended=True, Notify=False)
        self.ws = self.wb.Sheets("Formulario")
        self.ate_snap = LINHA_INI + TAM_SNAP - 1
        # estado original do modelo (linhas de exemplo), para restaurar depois
        self.original = self.ws.Range(f"{L_INI}{LINHA_INI}:{L_FIM}{self.ate_snap}").Value

    def fechar(self):
        try:
            self.wb.Close(SaveChanges=False)
        except Exception:
            pass

    def salvar(self, destino):
        self.wb.SaveCopyAs(os.path.abspath(destino))   # mantem o formato .xls do modelo

    def restaurar(self, escritos, ultima, destravado):
        """Devolve o modelo ao estado original e confere. False = nao reutilizar."""
        try:
            ws = self.ws
            if destravado:
                ws.Range(f"{L_INI}{LINHA_INI}:{L_FIM}{ultima}").ClearContents()
                k = min(ultima, self.ate_snap) - LINHA_INI + 1
                ws.Range(f"{L_INI}{LINHA_INI}:{L_FIM}{LINHA_INI + k - 1}").Value = self.original[:k]
            else:
                for j, r1, r2 in escritos:
                    letra = col_letra(COL_INI + j)
                    alvo = f"{letra}{r1}" if r1 == r2 else f"{letra}{r1}:{letra}{r2}"
                    ws.Range(alvo).ClearContents()
                    for r in range(r1, min(r2, self.ate_snap) + 1):
                        v = self.original[r - LINHA_INI][j]
                        if v is not None:
                            ws.Cells(r, COL_INI + j).Value = v

            fim = max(ultima, self.ate_snap)
            atual = ws.Range(f"{L_INI}{LINHA_INI}:{L_FIM}{fim}").Value
            esperado = tuple(self.original) + tuple((None,) * LARGURA for _ in range(fim - self.ate_snap))
            return atual == esperado
        except Exception:
            return False


def _descartar(cache):
    modelo = cache.get("modelo")
    if modelo is not None:
        modelo.fechar()
    cache["modelo"] = None


def processar_arquivo(excel, caminho, modelos, destino, cache):
    """Devolve (ok, mensagem, tipo_de_log)."""
    nome = os.path.basename(caminho)
    chave = identificar_modelo(nome)

    if not chave:
        return False, f"{nome}: modelo não identificado pelo nome", "erro"
    if chave not in MODELOS_VALIDOS:
        return False, f"{nome}: o modelo {chave} não existe", "erro"
    if chave not in modelos:
        return False, f"{nome}: modelo {chave} não foi carregado", "erro"

    try:
        linhas = obter_linhas(excel, caminho)
    except Exception as e:
        return False, f"{nome}: não foi possível ler ({_resumir_erro(e)})", "erro"
    if not linhas:
        return False, f"{nome}: nenhuma linha com dados", "aviso"

    novo_nome = os.path.splitext(nome)[0] + ".xls"
    try:
        modelo = cache.get("modelo")
        if modelo is None or modelo.chave != chave:
            _descartar(cache)
            modelo = cache["modelo"] = ModeloAberto(excel, modelos[chave], chave)
        escritos, destravado = gravar_modelo(modelo.ws, linhas)
        modelo.salvar(os.path.join(destino, novo_nome))
    except Exception as e:
        _descartar(cache)
        return False, f"{nome}: {_resumir_erro(e)}", "erro"

    ultima = LINHA_INI + len(linhas) - 1
    if not modelo.restaurar(escritos, ultima, destravado):
        _descartar(cache)   # proximo arquivo reabre o modelo limpo
    return True, novo_nome, "ok"


# ---------------- Excel e paralelismo ----------------
def criar_excel():
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AutomationSecurity = 3
    excel.EnableEvents = False
    try:
        excel.ScreenUpdating = False
    except Exception:
        pass
    return excel


def excel_vivo(excel):
    try:
        excel.Version
        return True
    except Exception:
        return False


def fechar_excel(excel):
    if excel is not None:
        try:
            excel.Quit()
        except Exception:
            pass


def executar_lote(arquivos, modelos, destino, n_excel, ui,
                  criar=criar_excel, processar=processar_arquivo):
    """Processa os arquivos com n_excel instancias do Excel em paralelo."""
    total = len(arquivos)
    pendentes = queue.Queue()
    for a in arquivos:
        pendentes.put(a)
    trava = threading.Lock()
    placar = {"feitos": 0, "ok": 0, "falhas": 0}

    def registrar(ok):
        with trava:
            placar["feitos"] += 1
            placar["ok" if ok else "falhas"] += 1
            feitos = placar["feitos"]
        ui("progresso", feitos / total, f"{feitos} / {total}")

    def trabalhador():
        pythoncom.CoInitialize()
        excel = None
        cache = {"modelo": None}
        try:
            while True:
                try:
                    caminho = pendentes.get_nowait()
                except queue.Empty:
                    return
                if excel is None:
                    try:
                        excel = criar()
                    except Exception as e:
                        ui("log", f"Não foi possível abrir o Excel: {_resumir_erro(e)}", "erro")
                        pendentes.put(caminho)
                        return
                ok, mensagem, tipo = processar(excel, caminho, modelos, destino, cache)
                ui("log", mensagem, tipo)
                registrar(ok)
                if not ok and not excel_vivo(excel):   # Excel caiu: recria no proximo arquivo
                    fechar_excel(excel)
                    excel = None
                    cache["modelo"] = None
        finally:
            _descartar(cache)
            fechar_excel(excel)
            pythoncom.CoUninitialize()

    threads = [threading.Thread(target=trabalhador, daemon=True) for _ in range(n_excel)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    while not pendentes.empty():   # nenhum Excel conseguiu abrir
        pendentes.get_nowait()
        registrar(False)
    return placar


# --------------------------------------------------------------------------
# Interface
# --------------------------------------------------------------------------
class AppTotumseg(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title(NOME_APP)
        self.geometry("640x560")
        self.minsize(580, 500)
        self.configure(fg_color=BRANCO)

        self.modelos = {}          # 'APTO_300' -> caminho do .xls
        self.pasta_origem = ""
        self.pasta_destino = ""
        self.fila = queue.Queue()  # comunicacao thread -> interface
        self.processando = False

        # cabecalho
        ctk.CTkLabel(self, text=NOME_APP, font=("Segoe UI", 20, "bold"),
                     text_color=AZUL).pack(anchor="w", padx=28, pady=(20, 12))
        ctk.CTkFrame(self, fg_color=AZUL, height=2, corner_radius=0).pack(fill="x")

        corpo = ctk.CTkFrame(self, fg_color=BRANCO)
        corpo.pack(fill="both", expand=True, padx=28, pady=(8, 20))
        corpo.grid_columnconfigure(0, weight=1)
        corpo.grid_rowconfigure(6, weight=1)

        self.lbl_modelos = self._linha(corpo, 0, "Modelos", "Carregar", self.selecionar_modelos)
        self.lbl_origem = self._linha(corpo, 1, "Pasta base", "Selecionar", self.selecionar_origem)
        self.lbl_destino = self._linha(corpo, 2, "Pasta destino", "Selecionar", self.selecionar_destino)

        self.btn_processar = ctk.CTkButton(
            corpo, text="Iniciar", command=self.iniciar_processo,
            font=("Segoe UI", 14, "bold"), height=42, corner_radius=8,
            fg_color=AZUL, hover_color=AZUL_ESCURO, text_color=BRANCO,
            text_color_disabled="#DBEAFE")
        self.btn_processar.grid(row=3, column=0, sticky="ew", pady=(16, 14))

        # progresso: uma linha de status + barra fina
        quadro = ctk.CTkFrame(corpo, fg_color=BRANCO)
        quadro.grid(row=4, column=0, sticky="ew")
        quadro.grid_columnconfigure(0, weight=1)
        self.lbl_status = ctk.CTkLabel(quadro, text="Aguardando", font=("Segoe UI", 12),
                                       text_color=TEXTO_SUAVE, anchor="w")
        self.lbl_status.grid(row=0, column=0, sticky="w")
        self.lbl_contagem = ctk.CTkLabel(quadro, text="", font=("Segoe UI", 12, "bold"),
                                         text_color=AZUL)
        self.lbl_contagem.grid(row=0, column=1, sticky="e")
        self.barra = ctk.CTkProgressBar(quadro, height=6, corner_radius=3,
                                        fg_color=AZUL_CLARO, progress_color=AZUL)
        self.barra.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(6, 0))
        self.barra.set(0)

        self.txt_log = ctk.CTkTextbox(corpo, fg_color=FUNDO_SUAVE, text_color=TEXTO,
                                      font=("Consolas", 11), corner_radius=8,
                                      border_width=1, border_color=BORDA, wrap="word")
        self.txt_log.grid(row=6, column=0, sticky="nsew", pady=(12, 0))
        for tag, cor in (("ok", VERDE), ("erro", VERMELHO), ("aviso", LARANJA)):
            self.txt_log.tag_config(tag, foreground=cor)
        self.txt_log.configure(state="disabled")

        self._carregar_config()
        self._atualizar_modelos()
        self._atualizar_pastas()
        if not TEM_EXCEL_COM:
            self.log("pywin32 não encontrado: é preciso Windows com Excel e pywin32 instalado.", "erro")
        self.after(100, self._ler_fila)

    # ---------------- construcao da tela ----------------
    def _linha(self, pai, linha, titulo, texto_botao, comando):
        """Uma linha simples: titulo + detalhe a esquerda, botao a direita."""
        quadro = ctk.CTkFrame(pai, fg_color=BRANCO)
        quadro.grid(row=linha, column=0, sticky="ew")
        quadro.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(quadro, text=titulo, font=("Segoe UI", 13, "bold"),
                     text_color=TEXTO, anchor="w").grid(row=0, column=0, sticky="sw", pady=(10, 0))
        detalhe = ctk.CTkLabel(quadro, text="", font=("Segoe UI", 11),
                               text_color=TEXTO_SUAVE, anchor="w")
        detalhe.grid(row=1, column=0, sticky="nw", pady=(0, 10))
        ctk.CTkButton(quadro, text=texto_botao, command=comando, width=110, height=32,
                      corner_radius=8, font=("Segoe UI", 12, "bold"),
                      fg_color=BRANCO, hover_color=AZUL_CLARO, text_color=AZUL,
                      border_width=1, border_color=AZUL).grid(row=0, column=1, rowspan=2)
        ctk.CTkFrame(quadro, fg_color=BORDA, height=1, corner_radius=0).grid(
            row=2, column=0, columnspan=2, sticky="ew")
        return detalhe

    # ---------------- estado da tela ----------------
    def _atualizar_modelos(self):
        n, total = len(self.modelos), len(MODELOS_VALIDOS)
        if n == 0:
            self.lbl_modelos.configure(text="Nenhum modelo carregado", text_color=TEXTO_SUAVE)
        elif n == total:
            self.lbl_modelos.configure(text=f"{n} de {total} carregados", text_color=VERDE)
        else:
            faltam = ", ".join(m for m in MODELOS_VALIDOS if m not in self.modelos)
            self.lbl_modelos.configure(text=encurtar(f"{n} de {total} carregados · faltam {faltam}", 70),
                                       text_color=AZUL)

    def _atualizar_pastas(self):
        for lbl, caminho in ((self.lbl_origem, self.pasta_origem), (self.lbl_destino, self.pasta_destino)):
            if caminho:
                lbl.configure(text=encurtar(caminho), text_color=TEXTO)
            else:
                lbl.configure(text="Não selecionada", text_color=TEXTO_SUAVE)

    def log(self, mensagem, tipo=None):
        prefixo = {"ok": "✓ ", "erro": "✗ ", "aviso": "! "}.get(tipo, "")
        self.txt_log.configure(state="normal")
        self.txt_log.insert("end", prefixo + mensagem + "\n", tipo)
        self.txt_log.see("end")
        self.txt_log.configure(state="disabled")

    # ---------------- selecoes ----------------
    def selecionar_modelos(self):
        arquivos = filedialog.askopenfilenames(
            title="Selecione os modelos (.xls) da Mapfre",
            filetypes=[("Planilhas Excel 97-2003", "*.xls")])
        if not arquivos:
            return
        recusados = []
        for caminho in arquivos:
            nome, ext = os.path.splitext(os.path.basename(caminho))
            if ext.lower() == ".xls" and nome.upper() in MODELOS_VALIDOS:
                self.modelos[nome.upper()] = caminho
            else:
                recusados.append(os.path.basename(caminho))
        self._atualizar_modelos()
        self._salvar_config()
        if recusados:
            messagebox.showwarning(
                "Alguns arquivos foram ignorados",
                "Os modelos precisam ter exatamente um destes nomes (.xls):\n\n"
                + ", ".join(MODELOS_VALIDOS)
                + "\n\nIgnorados:\n" + "\n".join(recusados))

    def selecionar_origem(self):
        pasta = filedialog.askdirectory(title="Selecione a pasta base (dados)")
        if pasta:
            self.pasta_origem = pasta
            self._atualizar_pastas()
            self._salvar_config()

    def selecionar_destino(self):
        pasta = filedialog.askdirectory(title="Selecione a pasta destino")
        if pasta:
            self.pasta_destino = pasta
            self._atualizar_pastas()
            self._salvar_config()

    # ---------------- configuracao salva ----------------
    def _salvar_config(self):
        try:
            with open(ARQ_CONFIG, "w", encoding="utf-8") as f:
                json.dump({"modelos": self.modelos, "origem": self.pasta_origem,
                           "destino": self.pasta_destino}, f, ensure_ascii=False)
        except OSError:
            pass

    def _carregar_config(self):
        try:
            with open(ARQ_CONFIG, encoding="utf-8") as f:
                cfg = json.load(f)
        except (OSError, ValueError):
            return
        self.modelos = {k: v for k, v in cfg.get("modelos", {}).items()
                        if k in MODELOS_VALIDOS and os.path.isfile(v)}
        if os.path.isdir(cfg.get("origem", "")):
            self.pasta_origem = cfg["origem"]
        if os.path.isdir(cfg.get("destino", "")):
            self.pasta_destino = cfg["destino"]

    # ---------------- execucao ----------------
    def iniciar_processo(self):
        if self.processando:
            return
        faltando = []
        if not self.modelos:
            faltando.append("carregar os modelos")
        if not self.pasta_origem:
            faltando.append("selecionar a pasta base")
        if not self.pasta_destino:
            faltando.append("selecionar a pasta destino")
        if faltando:
            messagebox.showwarning("Falta configurar", "Antes de iniciar:\n\n• " + "\n• ".join(faltando))
            return
        if not TEM_EXCEL_COM:
            messagebox.showerror("Excel não disponível",
                                 "Não foi possível usar o Excel (pywin32 não encontrado).")
            return

        self.processando = True
        self.btn_processar.configure(state="disabled", text="Processando…")
        self.barra.set(0)
        self.lbl_contagem.configure(text="")
        self.lbl_status.configure(text="Processando…")
        threading.Thread(
            target=self._processar,
            args=(dict(self.modelos), self.pasta_origem, self.pasta_destino),
            daemon=True).start()

    # A thread nunca mexe na interface direto: manda mensagens pela fila.
    def _ui(self, tipo, a=None, b=None):
        self.fila.put((tipo, a, b))

    def _ler_fila(self):
        try:
            while True:
                tipo, a, b = self.fila.get_nowait()
                if tipo == "log":
                    self.log(a, b)
                elif tipo == "status":
                    self.lbl_status.configure(text=a)
                elif tipo == "progresso":
                    self.barra.set(a)
                    self.lbl_contagem.configure(text=b)
                elif tipo == "fim":
                    self.processando = False
                    self.btn_processar.configure(state="normal", text="Iniciar")
                    self.lbl_status.configure(text=a)
        except queue.Empty:
            pass
        self.after(100, self._ler_fila)

    def _processar(self, modelos, origem, destino):
        inicio = time.perf_counter()
        resumo = "Concluído"
        try:
            arquivos = glob.glob(os.path.join(origem, "*.xlsm")) + glob.glob(os.path.join(origem, "*.xlsx"))
            arquivos = [a for a in arquivos if not os.path.basename(a).startswith("~$")]

            # dois arquivos com o mesmo nome-base gerariam o mesmo .xls
            vistos, unicos = set(), []
            for a in sorted(arquivos):
                chave = os.path.splitext(os.path.basename(a))[0].lower()
                if chave in vistos:
                    self._ui("log", f"{os.path.basename(a)}: ignorado (outro arquivo com o mesmo nome-base)", "aviso")
                else:
                    vistos.add(chave)
                    unicos.append(a)

            # agrupa por modelo: cada Excel reaproveita o modelo ja aberto
            arquivos = sorted(unicos, key=lambda a: (identificar_modelo(os.path.basename(a)) or "~", a))

            total = len(arquivos)
            if total == 0:
                self._ui("log", "Nenhum arquivo .xlsm/.xlsx na pasta base.", "aviso")
                resumo = "Nenhum arquivo encontrado"
                return

            n_excel = max(1, min(total, MAX_EXCEL, max(2, (os.cpu_count() or 2) - 1)))
            self._ui("status", "Abrindo o Excel…")
            self._ui("progresso", 0, f"0 / {total}")
            placar = executar_lote(arquivos, modelos, destino, n_excel, self._ui)

            segundos = f"{time.perf_counter() - inicio:.1f}".replace(".", ",")
            resumo = f"Concluído em {segundos} s"
            if placar["falhas"]:
                resumo += f" · {placar['falhas']} com problema"
            self._ui("log", f"{placar['ok']} de {total} gerado(s) em {segundos} s.",
                     "ok" if placar["falhas"] == 0 else "aviso")
        except Exception as e:
            self._ui("log", f"Erro fatal: {_resumir_erro(e)}", "erro")
            resumo = "Erro no processamento"
        finally:
            self._ui("fim", resumo)


if __name__ == "__main__":
    app = AppTotumseg()
    app.mainloop()
